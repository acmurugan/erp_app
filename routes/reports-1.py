from flask import Blueprint, render_template, request, redirect, url_for, session, flash, make_response, send_file
from datetime import datetime
import oracledb
from database import get_db_connection
import pandas as pd
import io
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.dataframe import dataframe_to_rows
from pathlib import Path

# Create the Blueprint
reports_bp = Blueprint('reports', __name__)

def get_report_config(report_id):
    """Enhanced version - Get report configuration from RPT_REPORT_MASTER with modular support"""
    try:
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                SELECT RPT_NAME, RPT_DESC, RPT_TYPE, RPT_SOURCE, RPT_PARAMS, RPT_OUTPUT_FORMATS
                FROM RPT_REPORT_MASTER 
                WHERE RPT_ID = :report_id 
                AND RPT_ACTIVE_FLAG = 'Y'
            """, {'report_id': report_id})
            
            result = cursor.fetchone()
            if result:
                # Handle LOB for RPT_SOURCE
                sql_query = result[3]
                if hasattr(sql_query, 'read'):
                    sql_query = sql_query.read()
                
                # NEW: Determine report type and if it's modular
                report_type = result[2] or 'SQL'  # Default to SQL if NULL
                
                # Check if this is a modular report (filename vs full SQL)
                is_modular = not (sql_query.strip().upper().startswith('SELECT') or 
                                sql_query.strip().upper().startswith('WITH') or
                                sql_query.strip().upper().startswith('--') or
                                len(sql_query.strip()) > 100)  # If it's short, likely a filename
                
                # Parse output formats
                output_formats_str = result[5] or 'VIEW,EXCEL'
                output_formats = [fmt.strip().lower() for fmt in output_formats_str.split(',')]
                
                print(f"📋 Report Type: {report_type}")
                print(f"📋 Is Modular: {is_modular}")
                print(f"📋 Source: {sql_query[:50]}{'...' if len(sql_query) > 50 else ''}")
                
                return {
                    'name': result[0],
                    'description': result[1],
                    'type': report_type,
                    'sql_query': sql_query,
                    'params': result[4],
                    'output_formats': output_formats,
                    'is_modular': is_modular  # NEW: Flag to identify modular reports
                }
            else:
                return None
    except Exception as e:
        print(f"❌ Error getting report config: {e}")
        return None

def get_report_sql(report_id):
    """Enhanced version - Get SQL query from database or external file"""
    try:
        config = get_report_config(report_id)
        if config:
            if config['is_modular']:
                # NEW: Load SQL from external file
                return load_sql_from_file(config['sql_query'], config['type'])
            else:
                # EXISTING: Return SQL directly from database
                return config['sql_query']
        else:
            raise Exception(f"Report {report_id} not found or inactive")
    except Exception as e:
        print(f"❌ Error getting report SQL: {e}")
        raise

def load_sql_from_file(source_reference, report_type):
    """NEW: Load SQL content from external file"""
    try:
        reports_base_path = Path("reports")
        
        if report_type == 'SQL':
            sql_file = reports_base_path / "sql" / f"{source_reference}.sql"
        elif report_type == 'TEMPLATE':
            sql_file = reports_base_path / "templates" / f"{source_reference}.sql"
        else:
            # For CLASS and PROCEDURE types, this function shouldn't be called
            raise ValueError(f"Cannot load SQL file for report type: {report_type}")
        
        if not sql_file.exists():
            raise FileNotFoundError(f"SQL file not found: {sql_file}")
        
        sql_content = sql_file.read_text(encoding='utf-8')
        print(f"📄 Loaded SQL from file: {sql_file}")
        print(f"📄 File size: {len(sql_content)} characters")
        return sql_content
        
    except Exception as e:
        print(f"❌ Error loading SQL from file: {e}")
        raise

def execute_report_with_params(report_id, params):
    """Enhanced version - Execute report with parameters supporting modular approach"""
    try:
        print(f"🔍 Executing enhanced report: {report_id}")
        print(f"📋 Parameters: {params}")
        
        # Get report configuration
        config = get_report_config(report_id)
        if not config:
            raise Exception(f"Report {report_id} not found")
        
        print(f"📊 Report Type: {config['type']}")
        print(f"🔧 Is Modular: {config['is_modular']}")
        
        # NEW: Handle different report types
        if config['type'] == 'CLASS':
            # Handle complex class-based reports
            return execute_class_report(report_id, config, params)
        elif config['type'] == 'PROCEDURE':
            # Handle stored procedure reports
            return execute_procedure_report(report_id, config, params)
        else:
            # Handle SQL and TEMPLATE reports (enhanced existing logic)
            return execute_sql_report(report_id, config, params)
            
    except Exception as e:
        print(f"❌ Error executing enhanced report: {e}")
        raise

def execute_sql_report(report_id, config, params):
    """Execute SQL report (both traditional and modular)"""
    try:
        # Get SQL query (either from database or file)
        if config['is_modular']:
            sql_query = load_sql_from_file(config['sql_query'], config['type'])
        else:
            sql_query = config['sql_query']
        
        # Your existing parameter replacement and execution logic
        query_length = len(sql_query) if sql_query else 0
        print(f"📋 SQL Query loaded, length: {query_length} characters")
        print(f"📋 Source: {'External File' if config['is_modular'] else 'Database'}")
        
        # Execute query (your existing logic)
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute(sql_query, params)
            
            # Get column names
            columns = [desc[0] for desc in cursor.description]
            
            # Fetch all results
            rows = cursor.fetchall()
            
            result = {
                'count': len(rows),
                'columns': columns,
                'data': rows,
                'query': str(sql_query)[:200] + "..." if len(str(sql_query)) > 200 else str(sql_query),
                'parameters': params,
                'report_id': report_id,
                'is_modular': config['is_modular'],  # NEW: Include modular flag
                'report_type': config['type']        # NEW: Include report type
            }
            
            print(f"✅ Query executed successfully. Rows: {len(rows)}")
            return result
            
    except Exception as e:
        print(f"❌ Error executing SQL report: {e}")
        raise

def execute_class_report(report_id, config, params):
    """NEW: Execute class-based complex reports"""
    try:
        print(f"🏗️ Executing class-based report: {config['sql_query']}")
        
        # Load and execute Python class
        reports_base_path = Path("reports")
        class_file = reports_base_path / "classes" / f"{config['sql_query']}.py"
        
        if not class_file.exists():
            # For now, return placeholder - implement full class loading later
            print(f"⚠️ Class file not found: {class_file}")
            return {
                'count': 1,
                'columns': ['Status', 'Message'],
                'data': [['INFO', f'Class file not found: {class_file}']],
                'query': f"Complex processing by {config['sql_query']}.py",
                'parameters': params,
                'report_id': report_id,
                'is_modular': True,
                'report_type': 'CLASS'
            }
        
        # TODO: Implement full class loading and execution
        # For now, return a placeholder
        return {
            'count': 1,
            'columns': ['Status', 'Message'],
            'data': [['SUCCESS', f'Complex report processor {config["sql_query"]} executed']],
            'query': f"Complex processing by {config['sql_query']}.py",
            'parameters': params,
            'report_id': report_id,
            'is_modular': True,
            'report_type': 'CLASS'
        }
        
    except Exception as e:
        print(f"❌ Error executing class report: {e}")
        raise

def execute_procedure_report(report_id, config, params):
    """NEW: Execute procedure-based reports"""
    try:
        print(f"⚙️ Executing procedure-based report: {config['sql_query']}")
        
        # TODO: Implement stored procedure execution
        # For now, return placeholder
        return {
            'count': 1,
            'columns': ['Status', 'Message'],
            'data': [['SUCCESS', f'Stored procedure {config["sql_query"]} executed']],
            'query': f"Procedure: {config['sql_query']}",
            'parameters': params,
            'report_id': report_id,
            'is_modular': True,
            'report_type': 'PROCEDURE'
        }
        
    except Exception as e:
        print(f"❌ Error executing procedure report: {e}")
        raise

def generate_excel_report(report_data, report_id, params):
    """Enhanced Excel generation with modular report support"""
    try:
        print(f"📊 Generating Excel report for {report_id}")
        
        # Create workbook and worksheet
        wb = Workbook()
        ws = wb.active
        ws.title = f"{report_id} Report"
        
        # Add title and metadata
        ws['A1'] = f"{report_id} - Report Results"
        ws['A1'].font = Font(bold=True, size=16)
        
        ws['A2'] = f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        ws['A3'] = f"Records: {report_data['count']}"
        
        # NEW: Add report type information
        if 'is_modular' in report_data:
            report_source = f"Source: {'External File' if report_data['is_modular'] else 'Database'}"
            if 'report_type' in report_data:
                report_source += f" ({report_data['report_type']})"
            ws['A4'] = report_source
            row_start = 6
        else:
            row_start = 5
        
        # Add parameters section
        row_num = row_start
        ws[f'A{row_num}'] = "Parameters:"
        ws[f'A{row_num}'].font = Font(bold=True)
        row_num += 1
        
        for key, value in params.items():
            if key not in ['comp_code', 'user_id']:  # Skip system parameters
                ws[f'A{row_num}'] = f"{key}: {value}"
                row_num += 1
        
        # Add data starting from row_num + 2
        data_start_row = row_num + 2
        
        # Add headers
        for col_num, column_name in enumerate(report_data['columns'], 1):
            cell = ws.cell(row=data_start_row, column=col_num, value=column_name)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # Add data rows
        for row_num, row_data in enumerate(report_data['data'], data_start_row + 1):
            for col_num, cell_value in enumerate(row_data, 1):
                # Handle None values and convert to string
                display_value = str(cell_value) if cell_value is not None else ""
                ws.cell(row=row_num, column=col_num, value=display_value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save to BytesIO
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        # Create response
        response = make_response(excel_buffer.getvalue())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename="{report_id}_Report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        
        print(f"✅ Excel file generated successfully")
        return response
        
    except Exception as e:
        print(f"❌ Error generating Excel: {e}")
        flash(f'Error generating Excel file: {str(e)}', 'danger')
        return redirect(request.referrer or '/dashboard')

def generate_pdf_report(report_data, report_id, params):
    """Enhanced PDF generation with modular report support"""
    try:
        print(f"📄 Generating PDF report for {report_id} with {len(report_data['columns'])} columns")
        
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter, A4, landscape, A3
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        
        # Create PDF buffer
        pdf_buffer = io.BytesIO()
        
        # Use A3 landscape for better wide table support
        doc = SimpleDocTemplate(pdf_buffer, pagesize=landscape(A3), 
                              leftMargin=0.3*inch, rightMargin=0.3*inch,
                              topMargin=0.3*inch, bottomMargin=0.3*inch)
        
        # Get styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=14,
            spaceAfter=20,
            alignment=1  # Center alignment
        )
        
        # Build story
        story = []
        
        # Add title
        title = Paragraph(f"{report_id} - Report Results", title_style)
        story.append(title)
        
        # Add metadata
        meta_style = styles['Normal']
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", meta_style))
        story.append(Paragraph(f"Records: {report_data['count']} | Columns: {len(report_data['columns'])}", meta_style))
        
        # NEW: Add report type information
        if 'is_modular' in report_data:
            report_source = f"Source: {'External File' if report_data['is_modular'] else 'Database'}"
            if 'report_type' in report_data:
                report_source += f" ({report_data['report_type']})"
            story.append(Paragraph(report_source, meta_style))
        
        story.append(Spacer(1, 12))
        
        # Add parameters
        story.append(Paragraph("Parameters:", styles['Heading2']))
        for key, value in params.items():
            if key not in ['comp_code', 'user_id']:
                story.append(Paragraph(f"{key}: {value}", meta_style))
        story.append(Spacer(1, 12))
        
        # Handle wide tables - split into multiple tables if necessary
        max_cols_per_table = 12  # Maximum columns per table for readability
        num_columns = len(report_data['columns'])
        
        if num_columns <= max_cols_per_table:
            # Single table for narrow reports
            print(f"📄 Creating single table with {num_columns} columns")
            table_data = [report_data['columns']]  # Headers
            max_rows = min(500, len(report_data['data']))  # Reduced for PDF performance
            
            for row in report_data['data'][:max_rows]:
                # Convert all values to strings and handle None
                pdf_row = [str(cell)[:25] + "..." if len(str(cell)) > 25 else str(cell) if cell is not None else "" for cell in row]
                table_data.append(pdf_row)
            
            # Create table
            table = Table(table_data)
            
            # Apply table style
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 6),
                ('FONTSIZE', (0, 1), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.black)
            ]))
            
            story.append(table)
        else:
            # Split into multiple tables for wide reports
            print(f"📄 Splitting {num_columns} columns into multiple tables")
            
            # Calculate number of table chunks needed
            num_chunks = (num_columns + max_cols_per_table - 1) // max_cols_per_table
            
            for chunk_index in range(num_chunks):
                start_col = chunk_index * max_cols_per_table
                end_col = min(start_col + max_cols_per_table, num_columns)
                
                # Add section header
                section_title = f"Columns {start_col + 1} to {end_col} (of {num_columns})"
                story.append(Paragraph(section_title, styles['Heading3']))
                story.append(Spacer(1, 6))
                
                # Create table for this chunk
                chunk_headers = report_data['columns'][start_col:end_col]
                table_data = [chunk_headers]
                
                max_rows = min(300, len(report_data['data']))  # Fewer rows per chunk
                
                for row in report_data['data'][:max_rows]:
                    chunk_row = row[start_col:end_col]
                    pdf_row = [str(cell)[:20] + "..." if len(str(cell)) > 20 else str(cell) if cell is not None else "" for cell in chunk_row]
                    table_data.append(pdf_row)
                
                # Create and style table
                table = Table(table_data)
                table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 7),
                    ('FONTSIZE', (0, 1), (-1, -1), 6),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.black)
                ]))
                
                story.append(table)
                
                # Add page break between chunks (except for the last one)
                if chunk_index < num_chunks - 1:
                    story.append(PageBreak())
        
        # Add notes
        story.append(Spacer(1, 12))
        notes = []
        
        if len(report_data['data']) > (500 if num_columns <= max_cols_per_table else 300):
            notes.append(f"Note: Only first {500 if num_columns <= max_cols_per_table else 300} records shown in PDF. Total records: {report_data['count']}")
        
        if num_columns > max_cols_per_table:
            notes.append(f"Note: {num_columns} columns split across multiple tables for better readability.")
            notes.append("For complete data in single view, use Excel export.")
        
        for note in notes:
            story.append(Paragraph(note, styles['Italic']))
        
        # Build PDF
        doc.build(story)
        pdf_buffer.seek(0)
        
        # Create response
        response = make_response(pdf_buffer.getvalue())
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'attachment; filename="{report_id}_Report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf"'
        
        print(f"✅ PDF file generated successfully with {num_columns} columns")
        return response
        
    except ImportError:
        print("❌ ReportLab not installed. Install with: pip install reportlab")
        flash('PDF generation requires ReportLab. Please contact administrator.', 'warning')
        return redirect(request.referrer or '/dashboard')
    except Exception as e:
        print(f"❌ Error generating PDF: {e}")
        flash(f'Error generating PDF file: {str(e)}', 'danger')
        return redirect(request.referrer or '/dashboard')

@reports_bp.route('/<form_id>/<menu_id>')
def report_parameter_form(form_id, menu_id):
    """Enhanced parameter form with modular report support"""
    if 'user_id' not in session:
        return redirect('/login')
    
    print(f"📊 Loading parameter form - Form ID: {form_id}, Menu ID: {menu_id}")
    
    # For TTL201 Customer Sales Analysis, get config from database
    if menu_id == 'T010509' or form_id == 'TTL201':
        # Get enhanced report configuration from database
        report_config = get_report_config('TTL201')
        
        if report_config:
            # Create form config with database values
            form_config = {
                'title': report_config['name'],
                'description': report_config['description'],
                'report_id': 'TTL201',
                'report_type': report_config['type'],          # NEW: Include report type
                'is_modular': report_config['is_modular'],     # NEW: Include modular flag
                'parameters': [
                    {
                        'name': 'SDM Code',
                        'field': 'sdm_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-user-tie'
                    },
                    {
                        'name': 'Customer Code',
                        'field': 'cust_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-users'
                    },
                    {
                        'name': 'Flash Code',
                        'field': 'flash_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-bolt'
                    },
                    {
                        'name': 'Item Code',
                        'field': 'item_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-box'
                    },
                    {
                        'name': 'Location Code',
                        'field': 'locn_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-map-marker-alt'
                    },
                    {
                        'name': 'Date',
                        'field': 'date',
                        'type': 'date_range',
                        'required': True,
                        'icon': 'fas fa-calendar-alt'
                    }
                ],
                'output_formats': report_config['output_formats']  # Dynamic from database
            }
        else:
            # Fallback config if database read fails
            form_config = {
                'title': 'Customerwise Sales Analysis',
                'description': 'Detailed sales analysis by customer with item breakdown, margins, and contribution analysis',
                'report_id': 'TTL201',
                'report_type': 'SQL',          # DEFAULT
                'is_modular': False,           # DEFAULT
                'parameters': [
                    {
                        'name': 'SDM Code',
                        'field': 'sdm_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-user-tie'
                    },
                    {
                        'name': 'Customer Code',
                        'field': 'cust_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-users'
                    },
                    {
                        'name': 'Flash Code',
                        'field': 'flash_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-bolt'
                    },
                    {
                        'name': 'Item Code',
                        'field': 'item_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-box'
                    },
                    {
                        'name': 'Location Code',
                        'field': 'locn_code',
                        'type': 'range',
                        'required': False,
                        'icon': 'fas fa-map-marker-alt'
                    },
                    {
                        'name': 'Date',
                        'field': 'date',
                        'type': 'date_range',
                        'required': True,
                        'icon': 'fas fa-calendar-alt'
                    }
                ],
                'output_formats': ['view', 'excel', 'pdf']  # Default fallback
            }
        
        response = render_template('reports/dynamic_parameter_form.html', 
                             form_config=form_config,
                             menu_id=menu_id,
                             today=datetime.now().strftime('%d/%m/%Y'))
        
        # Add headers to prevent caching and form resubmission
        response = make_response(response)
        response.headers['Cache-Control'] = 'no-cache, no-store, must-revalidate'
        response.headers['Pragma'] = 'no-cache'
        response.headers['Expires'] = '0'
        
        return response
    else:
        # Fallback for other reports
        return render_template('reports/parameter_form.html', 
                             form_id=form_id, 
                             menu_id=menu_id, 
                             menu_name='Report Parameters',
                             today=datetime.now().strftime('%d/%m/%Y'))

@reports_bp.route('/dynamic/<menu_id>/generate', methods=['POST'])
def generate_dynamic_report(menu_id):
    """Enhanced dynamic report generation with modular support"""
    if 'user_id' not in session:
        return redirect('/login')
    
    try:
        print(f"\n🔄 GENERATING ENHANCED DYNAMIC REPORT:")
        print(f"   Menu ID: {menu_id}")
        print(f"   User: {session['user_id']}")
        
        # Debug: Print all form data received
        print(f"📋 Raw form data:")
        for key, value in request.form.items():
            print(f"     {key}: '{value}'")
        
        # Get form metadata
        report_type = request.form.get('reportType', 'detail')
        output_format = request.form.get('outputFormat', 'view')
        report_id = request.form.get('report_id', 'TTL201')  # Default to TTL201
        
        print(f"   Report ID: {report_id}")
        print(f"   Report Type: {report_type}")
        print(f"   Output Format: {output_format}")
        
        # Build parameters for Oracle - SIMPLIFIED APPROACH
        def get_param_value(field_name, default_value=None):
            value = request.form.get(field_name, '').strip()
            if not value:
                return default_value
            return value
        
        # Map parameters exactly as SQL expects them
        oracle_params = {
            'comp_code': 'TTM',
            'user_id': session.get('user_id', '7010IP'),
            'from_sdm_code': get_param_value('from_sdm_code'),
            'to_sdm_code': get_param_value('to_sdm_code', 'ZZZZZZ'),
            'from_cust_code': get_param_value('from_cust_code'),
            'to_cust_code': get_param_value('to_cust_code', 'ZZZZZZ'),
            'from_flash_code': get_param_value('from_flash_code'),
            'to_flash_code': get_param_value('to_flash_code', 'ZZZZZZ'),
            'from_item_code': get_param_value('from_item_code'),
            'to_item_code': get_param_value('to_item_code', 'ZZZZZZ'),
            'from_locn_code': get_param_value('from_locn_code'),
            'to_locn_code': get_param_value('to_locn_code', 'ZZZZZZ'),
            'from_date': get_param_value('from_date', '01/01/2024'),
            'to_date': get_param_value('to_date', '31/12/2024')
        }
        
        print(f"🎯 Oracle Parameters:")
        for key, value in oracle_params.items():
            print(f"     :{key} = '{value}'")
        
        # Execute the enhanced report
        print("🔍 Executing enhanced report query...")
        report_data = execute_report_with_params(report_id, oracle_params)
        print(f"✅ Enhanced query executed successfully. Records: {report_data['count']}")
        
        # Handle different output formats
        if output_format == 'excel':
            print("📊 Generating Excel report...")
            return generate_excel_report(report_data, report_id, oracle_params)
        elif output_format == 'pdf':
            print("📄 Generating PDF report...")
            return generate_pdf_report(report_data, report_id, oracle_params)
        
        # Return results page for view format
        print("🖥️ Rendering enhanced report results...")
        return render_template('reports/report_results_simple.html',
                             data=report_data,
                             form_id=report_id,
                             menu_id=menu_id,
                             params=oracle_params,
                             generated_at=datetime.now().strftime('%d/%m/%Y %H:%M:%S'))
    
    except Exception as e:
        print(f"❌ Error generating enhanced dynamic report: {e}")
        import traceback
        traceback.print_exc()
        
        # Return error page with details
        error_info = {
            'error': str(e),
            'menu_id': menu_id,
            'form_data': dict(request.form),
            'traceback': traceback.format_exc()
        }
        
        return render_template('reports/error_simple.html', error_info=error_info)

@reports_bp.route('/<form_id>/<menu_id>/generate', methods=['POST'])
def generate_report(form_id, menu_id):
    """Enhanced generate report - SIMPLIFIED VERSION with modular support"""
    if 'user_id' not in session:
        return redirect('/login')
    
    try:
        print(f"\n🔄 GENERATING ENHANCED REPORT:")
        print(f"   Form ID: {form_id}")
        print(f"   Menu ID: {menu_id}")
        print(f"   User: {session['user_id']}")
        
        # Debug: Print all form data received
        print(f"📋 Raw form data:")
        for key, value in request.form.items():
            print(f"     {key}: '{value}'")
        
        # Redirect to dynamic report for TTL201
        if form_id == 'TTL201' or menu_id == 'T010509':
            return generate_dynamic_report(menu_id)
        
        # Handle other reports here if needed
        flash(f'Report {form_id} not implemented in this simplified version.', 'warning')
        return redirect('/dashboard')
        
    except Exception as e:
        print(f"❌ Error generating enhanced report: {e}")
        import traceback
        traceback.print_exc()
        
        flash(f'Error generating report: {str(e)}', 'danger')
        return redirect('/reports/{form_id}/{menu_id}')

# NEW: Admin routes to manage modular reports
@reports_bp.route('/admin/modular')
def admin_modular_reports():
    """Simplified admin page for debugging"""
    if 'user_id' not in session:
        return redirect('/login')
    
    try:
        print("🔍 Starting admin_modular_reports route")
        reports_info = []
        
        print("🔍 Attempting database connection...")
        with get_db_connection() as conn:
            cursor = conn.cursor()
            print("🔍 Executing query...")
            cursor.execute("""
                SELECT RPT_ID, RPT_NAME, RPT_TYPE, RPT_SOURCE, RPT_ACTIVE_FLAG
                FROM RPT_REPORT_MASTER 
                WHERE RPT_ACTIVE_FLAG = 'Y'
                ORDER BY RPT_ID
            """)
            
            print("🔍 Processing results...")
            row_count = 0
            for row in cursor.fetchall():
                row_count += 1
                report_id, name, report_type, source, active = row
                
                print(f"🔍 Processing row {row_count}: {report_id}")
                
                # Handle LOB safely
                try:
                    if hasattr(source, 'read'):
                        source_content = source.read()
                    else:
                        source_content = str(source) if source else ""
                except Exception as e:
                    print(f"⚠️ Error reading source for {report_id}: {e}")
                    source_content = "Error reading source"
                
                # Simple modular detection
                try:
                    source_upper = source_content.strip().upper()
                    is_modular = not (source_upper.startswith('SELECT') or 
                                    source_upper.startswith('WITH') or
                                    len(source_content.strip()) > 100)
                except:
                    is_modular = False
                
                # Simple file existence check
                file_exists = True
                file_path = ""
                try:
                    if is_modular and (report_type == 'SQL' or not report_type):
                        file_path = f"reports/sql/{source_content}.sql"
                        file_exists = Path(file_path).exists()
                except Exception as e:
                    print(f"⚠️ Error checking file for {report_id}: {e}")
                    file_exists = False
                
                # Determine status
                if is_modular and file_exists:
                    status = 'Modular File OK'
                elif is_modular and not file_exists:
                    status = 'Missing File'
                else:
                    status = 'Traditional SQL'
                
                reports_info.append({
                    'report_id': report_id,
                    'name': name or 'No Name',
                    'type': report_type or 'SQL',
                    'source': source_content[:50] + "..." if len(source_content) > 50 else source_content,
                    'is_modular': is_modular,
                    'file_exists': file_exists,
                    'file_path': file_path,
                    'status': status,
                    'created_date': None,  # Simplified for debugging
                    'updated_date': None   # Simplified for debugging
                })
            
            print(f"✅ Processed {row_count} reports successfully")
        
        # Try to render a simple template first
        print("🔍 Attempting to render template...")
        
        # Simple HTML response for debugging
        html_response = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Debug Admin Page</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>🔍 Debug Admin Page</h2>
                <p><strong>Reports Found:</strong> {len(reports_info)}</p>
                
                <div class="alert alert-success">
                    ✅ Database connection and query successful!
                </div>
                
                <table class="table table-striped">
                    <thead>
                        <tr>
                            <th>Report ID</th>
                            <th>Name</th>
                            <th>Type</th>
                            <th>Status</th>
                            <th>Actions</th>
                        </tr>
                    </thead>
                    <tbody>
        """
        
        for report in reports_info:
            migrate_button = ""
            if report['status'] == 'Traditional SQL':
                migrate_button = f'''
                <a href="/reports/admin/migrate/{report['report_id']}" 
                   class="btn btn-sm btn-info"
                   onclick="return confirm('Migrate {report['report_id']} to modular approach?')">
                    🔄 Migrate
                </a>
                '''
            
            html_response += f"""
                        <tr>
                            <td><code>{report['report_id']}</code></td>
                            <td>{report['name']}</td>
                            <td><span class="badge bg-secondary">{report['type']}</span></td>
                            <td>
                                {"✅" if report['status'] == 'Modular File OK' else 
                                 "⚠️" if report['status'] == 'Missing File' else "📋"} 
                                {report['status']}
                            </td>
                            <td>
                                <a href="/reports/{report['report_id']}/ADMIN" class="btn btn-sm btn-primary">🔵 Test</a>
                                <a href="/reports/debug/modular/{report['report_id']}" class="btn btn-sm btn-warning">🐛 Debug</a>
                                {migrate_button}
                            </td>
                        </tr>
            """
        
        html_response += """
                    </tbody>
                </table>
                
                <div class="mt-4">
                    <a href="/dashboard" class="btn btn-outline-secondary">← Back to Dashboard</a>
                    <a href="/reports/debug/test" class="btn btn-outline-info">🧪 Debug Test</a>
                </div>
            </div>
        </body>
        </html>
        """
        
        print("✅ Returning HTML response")
        return html_response
        
    except Exception as e:
        print(f"❌ Error in admin_modular_reports: {e}")
        import traceback
        error_trace = traceback.format_exc()
        print(f"❌ Full traceback:\n{error_trace}")
        
        # Return error information
        error_html = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Admin Error Debug</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>❌ Admin Route Error</h2>
                <div class="alert alert-danger">
                    <strong>Error:</strong> {str(e)}
                </div>
                <details>
                    <summary>Full Traceback</summary>
                    <pre class="bg-light p-3">{error_trace}</pre>
                </details>
                <div class="mt-4">
                    <a href="/dashboard" class="btn btn-outline-secondary">← Back to Dashboard</a>
                    <a href="/reports/debug/test" class="btn btn-outline-info">🧪 Debug Test</a>
                </div>
            </div>
        </body>
        </html>
        """
        return error_html

@reports_bp.route('/debug/admin-test')
def debug_admin_test():
    """Simple test to check if admin routes work"""
    return f"""
    <!DOCTYPE html>
    <html>
    <head>
        <title>Admin Debug Test</title>
        <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
    </head>
    <body>
        <div class="container mt-4">
            <h2>🧪 Admin Debug Test</h2>
            <div class="alert alert-success">
                ✅ Reports blueprint is working!
            </div>
            <p><strong>Session User:</strong> {session.get('user_id', 'Not logged in')}</p>
            <p><strong>Time:</strong> {datetime.now()}</p>
            <div class="mt-3">
                <a href="/reports/admin/modular" class="btn btn-primary">→ Try Admin Modular</a>
                <a href="/dashboard" class="btn btn-secondary">← Back to Dashboard</a>
            </div>
        </div>
    </body>
    </html>
    """

@reports_bp.route('/debug/test', methods=['GET', 'POST'])
def debug_test():
    """Simple debug test route"""
    if request.method == 'POST':
        return f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Debug Test Results</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>Debug Test - POST Request Received</h2>
                <h3>Form Data:</h3>
                <pre class="bg-light p-3">{dict(request.form)}</pre>
                <h3>Session:</h3>
                <pre class="bg-light p-3">{dict(session)}</pre>
                <a href="/dashboard" class="btn btn-secondary">Back to Dashboard</a>
            </div>
        </body>
        </html>
        """
    else:
        return f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Debug Test Form</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>Debug Test - GET Request</h2>
                <form method="POST" class="mt-3">
                    <div class="mb-3">
                        <label class="form-label">Test Field:</label>
                        <input type="text" name="test_field" value="test_value" class="form-control">
                    </div>
                    <button type="submit" class="btn btn-primary">Test Submit</button>
                    <a href="/dashboard" class="btn btn-secondary">Back to Dashboard</a>
                </form>
            </div>
        </body>
        </html>
        """

@reports_bp.route('/debug/params', methods=['GET', 'POST'])
def debug_parameters():
    """Debug route to check parameter mapping"""
    if 'user_id' not in session:
        return redirect('/login')  # Use direct path instead of url_for
    
    if request.method == 'POST':
        print("🐛 DEBUG PARAMETERS ROUTE")
        print("📋 Raw Form Data:")
        for key, value in request.form.items():
            print(f"   {key}: '{value}'")
        
        return f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Parameter Debug Results</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>Parameter Mapping Debug</h2>
                <h3>Raw Form Data:</h3>
                <pre class="bg-light p-3">{dict(request.form)}</pre>
                <a href="/dashboard" class="btn btn-secondary">Back to Dashboard</a>
            </div>
        </body>
        </html>
        """
    else:
        # GET request - show a test form
        return f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Parameter Debug Form</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>Test Parameter Form</h2>
                <form method="POST">
                    <div class="row">
                        <div class="col-md-6">
                            <label>From Customer Code:</label>
                            <input type="text" name="from_cust_code" value="0" class="form-control">
                        </div>
                        <div class="col-md-6">
                            <label>To Customer Code:</label>
                            <input type="text" name="to_cust_code" value="ZZZZZZ" class="form-control">
                        </div>
                    </div>
                    <div class="row mt-3">
                        <div class="col-md-6">
                            <label>From Date:</label>
                            <input type="text" name="from_date" value="01/01/2024" class="form-control">
                        </div>
                        <div class="col-md-6">
                            <label>To Date:</label>
                            <input type="text" name="to_date" value="31/12/2024" class="form-control">
                        </div>
                    </div>
                    <div class="mt-3">
                        <button type="submit" class="btn btn-primary">Test Mapping</button>
                        <a href="/dashboard" class="btn btn-secondary">Back to Dashboard</a>
                    </div>
                </form>
            </div>
        </body>
        </html>
        """

@reports_bp.route('/debug/modular/<report_id>')
def debug_modular_report(report_id):
    """NEW: Debug route to check modular report configuration"""
    if 'user_id' not in session:
        return redirect('/login')  # Use direct path instead of url_for
    
    try:
        config = get_report_config(report_id)
        
        if not config:
            return f"""
            <!DOCTYPE html>
            <html>
            <head>
                <title>Report Not Found</title>
                <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
            </head>
            <body>
                <div class="container mt-4">
                    <h2>Report {report_id} not found</h2>
                    <a href="/reports/admin/modular" class="btn btn-secondary">Back to Admin</a>
                </div>
            </body>
            </html>
            """, 404
        
        debug_info = {
            'report_id': report_id,
            'config': config,
            'file_check': {
                'is_modular': config['is_modular'],
                'type': config['type'],
                'source': config['sql_query']
            }
        }
        
        # Check file existence if modular
        if config['is_modular']:
            if config['type'] == 'SQL' or not config['type']:
                expected_path = Path("reports/sql") / f"{config['sql_query']}.sql"
            elif config['type'] == 'CLASS':
                expected_path = Path("reports/classes") / f"{config['sql_query']}.py"
            elif config['type'] == 'TEMPLATE':
                expected_path = Path("reports/templates") / f"{config['sql_query']}.sql"
            else:
                expected_path = Path("reports/procedures") / f"{config['sql_query']}.json"
            
            debug_info['file_check']['expected_path'] = str(expected_path)
            debug_info['file_check']['file_exists'] = expected_path.exists()
            
            if expected_path.exists():
                try:
                    debug_info['file_check']['file_content'] = expected_path.read_text(encoding='utf-8')[:500] + "..."
                except:
                    debug_info['file_check']['file_content'] = "Error reading file"
        
        return f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Debug Report: {report_id}</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>Debug Report: {report_id}</h2>
                <h3>Configuration:</h3>
                <pre class="bg-light p-3">{json.dumps(debug_info, indent=2, default=str)}</pre>
                <div class="mt-3">
                    <a href="/reports/admin/modular" class="btn btn-secondary">Back to Admin</a>
                    <a href="/reports/{report_id}/TEST" class="btn btn-primary">Test Report</a>
                </div>
            </div>
        </body>
        </html>
        """
        
    except Exception as e:
        import traceback
        return f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>Debug Error</title>
            <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
        </head>
        <body>
            <div class="container mt-4">
                <h2>Debug Error</h2>
                <div class="alert alert-danger">
                    <strong>Error:</strong> {str(e)}
                </div>
                <details>
                    <summary>Full Traceback</summary>
                    <pre class="bg-light p-3">{traceback.format_exc()}</pre>
                </details>
                <a href="/reports/admin/modular" class="btn btn-secondary">Back to Admin</a>
            </div>
        </body>
        </html>
        """

# NEW: Migration route for converting traditional reports to modular
@reports_bp.route('/admin/migrate/<report_id>')
def migrate_report_to_modular(report_id):
    """Convert a traditional SQL report to modular approach"""
    if 'user_id' not in session:
        return redirect('/login')
    
    try:
        config = get_report_config(report_id)
        if not config:
            flash(f'Report {report_id} not found', 'error')
            return redirect('/reports/admin/modular')
        
        if config['is_modular']:
            flash(f'Report {report_id} is already modular', 'info')
            return redirect('/reports/admin/modular')
        
        # Create SQL file from database content
        sql_content = config['sql_query']
        reports_sql_dir = Path("reports/sql")
        reports_sql_dir.mkdir(parents=True, exist_ok=True)
        
        sql_file_path = reports_sql_dir / f"{report_id}.sql"
        sql_file_path.write_text(sql_content, encoding='utf-8')
        
        # Update database to point to file instead of containing SQL
        with get_db_connection() as conn:
            cursor = conn.cursor()
            cursor.execute("""
                UPDATE RPT_REPORT_MASTER 
                SET RPT_SOURCE = :filename,
                    RPT_TYPE = 'SQL'
                WHERE RPT_ID = :report_id
            """, {'filename': report_id, 'report_id': report_id})
            conn.commit()
        
        flash(f'Report {report_id} successfully migrated to modular approach', 'success')
        return redirect('/reports/admin/modular')
        
    except Exception as e:
        print(f"❌ Error migrating report {report_id}: {e}")
        flash(f'Error migrating report {report_id}: {str(e)}', 'error')
        return redirect('/reports/admin/modular')