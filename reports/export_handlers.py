# =============================================================================
# 6. reports/export_handlers.py - Excel and PDF Generation
# =============================================================================

from flask import make_response, flash, redirect, request
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

def generate_excel_report(report_data, report_id, params):
    """Generate Excel report"""
    try:
        print(f"CHART/INFO Generating Excel report for {report_id}")
        
        wb = Workbook()
        ws = wb.active
        ws.title = f"{report_id} Report"
        
        # Add title and metadata
        ws['A1'] = f"{report_id} - Report Results"
        ws['A1'].font = Font(bold=True, size=16)
        
        ws['A2'] = f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        ws['A3'] = f"Records: {report_data['count']}"
        
        # Add report type information
        if 'is_modular' in report_data:
            report_source = f"Source: {'External File' if report_data['is_modular'] else 'Database'}"
            if 'report_type' in report_data:
                report_source += f" ({report_data['report_type']})"
            ws['A4'] = report_source
            row_start = 6
        else:
            row_start = 5
        
        # Add parameters
        row_num = row_start
        ws[f'A{row_num}'] = "Parameters:"
        ws[f'A{row_num}'].font = Font(bold=True)
        row_num += 1
        
        for key, value in params.items():
            if key not in ['comp_code', 'user_id']:
                ws[f'A{row_num}'] = f"{key}: {value}"
                row_num += 1
        
        # Add data
        data_start_row = row_num + 2
        
        # Headers
        for col_num, column_name in enumerate(report_data['columns'], 1):
            cell = ws.cell(row=data_start_row, column=col_num, value=column_name)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
        
        # Data rows
        for row_num, row_data in enumerate(report_data['data'], data_start_row + 1):
            for col_num, cell_value in enumerate(row_data, 1):
                display_value = str(cell_value) if cell_value is not None else ""
                ws.cell(row=row_num, column=col_num, value=display_value)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Save to BytesIO
        excel_buffer = io.BytesIO()
        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        # Create response
        response = make_response(excel_buffer.getvalue())
        response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        response.headers['Content-Disposition'] = f'attachment; filename="{report_id}_Report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'
        
        print(f"SUCCESS Excel file generated successfully")
        return response
        
    except Exception as e:
        print(f"ERROR Error generating Excel: {e}")
        flash(f'Error generating Excel file: {str(e)}', 'danger')
        return redirect(request.referrer or '/dashboard')

def generate_pdf_report(report_data, report_id, params):
    """Generate PDF report"""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, A3
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import inch
        
        pdf_buffer = io.BytesIO()
        doc = SimpleDocTemplate(pdf_buffer, pagesize=landscape(A3), 
                              leftMargin=0.3*inch, rightMargin=0.3*inch,
                              topMargin=0.3*inch, bottomMargin=0.3*inch)
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'],
                                    fontSize=14, spaceAfter=20, alignment=1)
        
        story = []
        
        # Add title and metadata
        story.append(Paragraph(f"{report_id} - Report Results", title_style))
        story.append(Paragraph(f"Generated: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}", styles['Normal']))
        story.append(Paragraph(f"Records: {report_data['count']} | Columns: {len(report_data['columns'])}", styles['Normal']))
        
        if 'is_modular' in report_data:
            report_source = f"Source: {'External File' if report_data['is_modular'] else 'Database'}"
            if 'report_type' in report_data:
                report_source += f" ({report_data['report_type']})"
            story.append(Paragraph(report_source, styles['Normal']))
        
        story.append(Spacer(1, 12))
        
        # Add parameters
        story.append(Paragraph("Parameters:", styles['Heading2']))
        for key, value in params.items():
            if key not in ['comp_code', 'user_id']:
                story.append(Paragraph(f"{key}: {value}", styles['Normal']))
        story.append(Spacer(1, 12))
        
        # Create table data
        table_data = [report_data['columns']]
        max_rows = min(500, len(report_data['data']))
        
        for row in report_data['data'][:max_rows]:
            pdf_row = [str(cell)[:25] + "..." if len(str(cell)) > 25 else str(cell) if cell is not None else "" for cell in row]
            table_data.append(pdf_row)
        
        # Create and style table
        table = Table(table_data)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 6),
            ('FONTSIZE', (0, 1), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.black)
        ]))
        
        story.append(table)
        
        # Build PDF
        doc.build(story)
        pdf_buffer.seek(0)
        
        # Create response
        response = make_response(pdf_buffer.getvalue())
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'attachment; filename="{report_id}_Report_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf"'
        
        return response
        
    except ImportError:
        flash('PDF generation requires ReportLab. Please contact administrator.', 'warning')
        return redirect(request.referrer or '/dashboard')
    except Exception as e:
        flash(f'Error generating PDF file: {str(e)}', 'danger')
        return redirect(request.referrer or '/dashboard')
